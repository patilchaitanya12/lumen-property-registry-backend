from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select

from app.db.session import SessionLocal
from app.models.owner import Owner


BATCH_SIZE = 1000

SOURCE_COLUMNS = [
    "owner_id",
    "first_name",
    "last_name",
    "title",
    "vip_tier",
    "source_community",
    "source_file",
    "gender",
    "gender_source",
    "property_count",
    "communities_owned",
    "community_list",
    "buildings_list",
    "is_multi_property",
    "is_portfolio_investor",
    "has_cross_community",
    "portfolio_tier",
    "has_plot",
    "has_apartment",
    "has_villa",
    "is_reachable",
]


def clean_string(value):
    if value is None:
        return None

    value = str(value).strip()

    return value if value else None


def clean_int(value):
    value = clean_string(value)

    if value is None:
        return None

    try:
        return int(value)
    except ValueError:
        raise ValueError(f"Invalid integer value: {value!r}")


def clean_bool(value):
    value = clean_string(value)

    if value is None:
        return None

    normalized = value.lower()

    if normalized == "true":
        return True

    if normalized == "false":
        return False

    raise ValueError(f"Invalid boolean value: {value!r}")


def build_owner_values(row, indexes):
    return {
        "first_name": clean_string(row[indexes["first_name"]]),
        "last_name": clean_string(row[indexes["last_name"]]),
        "title": clean_string(row[indexes["title"]]),
        "vip_tier": clean_string(row[indexes["vip_tier"]]),
        "source_community": clean_string(
            row[indexes["source_community"]]
        ),
        "source_file": clean_string(row[indexes["source_file"]]),
        "gender": clean_string(row[indexes["gender"]]),
        "gender_source": clean_string(
            row[indexes["gender_source"]]
        ),
        "property_count": clean_int(
            row[indexes["property_count"]]
        ),
        "communities_owned": clean_int(
            row[indexes["communities_owned"]]
        ),
        "community_list": clean_string(
            row[indexes["community_list"]]
        ),
        "buildings_list": clean_string(
            row[indexes["buildings_list"]]
        ),
        "is_multi_property": clean_bool(
            row[indexes["is_multi_property"]]
        ),
        "is_portfolio_investor": clean_bool(
            row[indexes["is_portfolio_investor"]]
        ),
        "has_cross_community": clean_bool(
            row[indexes["has_cross_community"]]
        ),
        "portfolio_tier": clean_string(
            row[indexes["portfolio_tier"]]
        ),
        "has_plot": clean_bool(row[indexes["has_plot"]]),
        "has_apartment": clean_bool(
            row[indexes["has_apartment"]]
        ),
        "has_villa": clean_bool(row[indexes["has_villa"]]),
        "is_reachable": clean_bool(
            row[indexes["is_reachable"]]
        ),
    }


def main():
    if len(sys.argv) != 2:
        print(
            "Usage: uv run python scripts/backfill_owner_profile.py "
            "<xlsx_path>"
        )
        raise SystemExit(1)

    xlsx_path = Path(sys.argv[1])

    if not xlsx_path.exists():
        print(f"File not found: {xlsx_path}")
        raise SystemExit(1)

    print(f"Reading: {xlsx_path}")

    wb = load_workbook(
        xlsx_path,
        read_only=True,
        data_only=True,
    )

    ws = wb["OWNERS"]

    rows = ws.iter_rows(values_only=True)
    headers = list(next(rows))

    missing_columns = [
        column
        for column in SOURCE_COLUMNS
        if column not in headers
    ]

    if missing_columns:
        wb.close()
        raise RuntimeError(
            f"Missing expected columns: {missing_columns}"
        )

    indexes = {
        column: headers.index(column)
        for column in SOURCE_COLUMNS
    }

    db = SessionLocal()

    try:
        total_rows = 0
        updated = 0
        not_found = 0
        batch = []

        for row in rows:
            total_rows += 1

            owner_id = clean_string(
                row[indexes["owner_id"]]
            )

            if not owner_id:
                continue

            batch.append(
                (
                    owner_id,
                    build_owner_values(row, indexes),
                )
            )

            if len(batch) >= BATCH_SIZE:
                updated_batch, not_found_batch = process_batch(
                    db,
                    batch,
                )

                updated += updated_batch
                not_found += not_found_batch
                batch.clear()

                if total_rows % 10000 == 0:
                    print(
                        f"Processed: {total_rows:,} | "
                        f"Updated: {updated:,} | "
                        f"Not found: {not_found:,}"
                    )

        if batch:
            updated_batch, not_found_batch = process_batch(
                db,
                batch,
            )

            updated += updated_batch
            not_found += not_found_batch

        print()
        print("Backfill complete")
        print(f"Rows processed: {total_rows:,}")
        print(f"Owners updated: {updated:,}")
        print(f"Owners not found: {not_found:,}")

    except Exception:
        db.rollback()
        raise

    finally:
        db.close()
        wb.close()


def process_batch(db, batch):
    owner_ids = [owner_id for owner_id, _ in batch]

    owners = db.scalars(
        select(Owner).where(
            Owner.owner_id.in_(owner_ids)
        )
    ).all()

    owners_by_id = {
        owner.owner_id: owner
        for owner in owners
    }

    updated = 0
    not_found = 0

    for owner_id, values in batch:
        owner = owners_by_id.get(owner_id)

        if owner is None:
            not_found += 1
            continue

        for field, value in values.items():
            setattr(owner, field, value)

        updated += 1

    db.commit()

    return updated, not_found


if __name__ == "__main__":
    main()
