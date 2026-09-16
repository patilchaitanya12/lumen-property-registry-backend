from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook


DEFAULT_FILE = Path("data/raw/04_data_model_v7.xlsx")

SOURCE_SHEETS = {
    "owners": "OWNERS",
    "properties": "PROPERTIES",
    "properties_new": "PROPERTIES_NEW",
    "contacts": "CONTACTS",
}


def clean_value(value):
    """Convert workbook placeholders to usable Python values."""
    if value is None:
        return None

    if isinstance(value, str):
        value = value.strip()

        if not value or value.upper() in {"[NOT FOUND]", "N/A", "NA", "NULL"}:
            return None

        return value

    return value


def normalize_bool(value):
    """Convert Excel boolean-like values to Python bool."""
    value = clean_value(value)

    if value is None:
        return None

    if isinstance(value, bool):
        return value

    if isinstance(value, str):
        normalized = value.strip().lower()

        if normalized in {"true", "1", "yes", "y"}:
            return True

        if normalized in {"false", "0", "no", "n"}:
            return False

    return None


def read_sheet(path: Path, sheet_name: str, limit: int | None = None):
    """Stream rows from an Excel sheet without loading the entire sheet."""
    workbook = load_workbook(
        path,
        read_only=True,
        data_only=True,
    )

    try:
        worksheet = workbook[sheet_name]
        rows = worksheet.iter_rows(values_only=True)

        headers = next(rows)
        headers = [str(header).strip() if header else None for header in headers]

        for row_number, row in enumerate(rows, start=2):
            if limit is not None and row_number > limit + 1:
                break

            yield row_number, dict(zip(headers, row))
    finally:
        workbook.close()


def inspect_sheet(path: Path, sheet_name: str, limit: int = 3):
    """Print headers and a few cleaned rows for validation."""
    print(f"\n{'=' * 80}")
    print(f"SHEET: {sheet_name}")
    print("=" * 80)

    for row_number, row in read_sheet(path, sheet_name, limit):
        cleaned = {
            key: clean_value(value)
            for key, value in row.items()
            if key is not None
        }

        print(f"Row {row_number}: {cleaned}")


def validate_properties(path: Path, limit: int | None = None):
    """
    Validate property identity across PROPERTIES and PROPERTIES_NEW.

    We expect:
    - property_id is present
    - PROPERTIES may contain duplicate property IDs
    - PROPERTIES_NEW should contain unique property IDs
    - the two sheets should not overlap
    """
    print("\nValidating property identity...")

    property_ids: set[str] = set()
    properties_duplicates: set[str] = set()
    properties_new_ids: set[str] = set()

    for row_number, row in read_sheet(
        path,
        SOURCE_SHEETS["properties"],
        limit,
    ):
        property_id = clean_value(row.get("property_id"))

        if property_id is None:
            raise ValueError(
                f"Missing property_id in PROPERTIES row {row_number}"
            )

        property_id = str(property_id)

        if property_id in property_ids:
            properties_duplicates.add(property_id)

        property_ids.add(property_id)

    for row_number, row in read_sheet(
        path,
        SOURCE_SHEETS["properties_new"],
        limit,
    ):
        property_id = clean_value(row.get("property_id"))

        if property_id is None:
            raise ValueError(
                f"Missing property_id in PROPERTIES_NEW row {row_number}"
            )

        property_id = str(property_id)

        if property_id in properties_new_ids:
            raise ValueError(
                f"Duplicate property_id in PROPERTIES_NEW: {property_id}"
            )

        properties_new_ids.add(property_id)

    overlap = property_ids & properties_new_ids

    print(f"PROPERTIES unique IDs:     {len(property_ids):,}")
    print(f"PROPERTIES duplicate IDs:  {len(properties_duplicates):,}")
    print(f"PROPERTIES_NEW unique IDs: {len(properties_new_ids):,}")
    print(f"Cross-sheet overlap:       {len(overlap):,}")

    if overlap:
        sample = list(overlap)[:10]
        raise ValueError(
            f"Property IDs appear in both sheets: {sample}"
        )


def run_dry_run(path: Path, limit: int):
    """Run validation without writing anything to PostgreSQL."""
    print("Lumen Property Registry - ingestion dry run")
    print(f"Source: {path}")
    print(f"Row limit per large sheet: {limit:,}")

    if not path.exists():
        raise FileNotFoundError(f"Source file not found: {path}")

    for sheet_name in SOURCE_SHEETS.values():
        inspect_sheet(path, sheet_name, limit=3)

    validate_properties(path, limit=limit)

    print("\nDry run completed successfully.")
    print("No database changes were made.")


def main():
    parser = argparse.ArgumentParser(
        description="Ingest the Lumen processed Excel dataset."
    )

    parser.add_argument(
        "--file",
        type=Path,
        default=DEFAULT_FILE,
        help="Path to the processed Excel workbook.",
    )

    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Validate the source without writing to the database.",
    )

    parser.add_argument(
        "--limit",
        type=int,
        default=1000,
        help="Rows to inspect from each large sheet during dry run.",
    )

    args = parser.parse_args()

    if args.dry_run:
        run_dry_run(args.file, args.limit)
        return

    raise NotImplementedError(
        "Full database ingestion will be enabled after dry-run validation."
    )


if __name__ == "__main__":
    main()